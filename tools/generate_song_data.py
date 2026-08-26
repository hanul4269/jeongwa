import json
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path("/Users/hanul/outputs/jeongwa_songbook_screenshot_20260811/jeongwa_songbook_visible_rows.xlsx")
OUTPUT = ROOT / "songs-data.js"


def clean(value):
    if value is None:
        return ""
    return str(value).strip()


wb = load_workbook(SOURCE, data_only=True)
ws = wb["노래책"]

headers = [clean(cell.value) for cell in ws[1]]


def find_col(*names):
    normalized = {name.lower(): idx for idx, name in enumerate(headers)}
    for name in names:
        idx = normalized.get(name.lower())
        if idx is not None:
            return idx
    return None


def row_value(row, *names):
    idx = find_col(*names)
    if idx is None or idx >= len(row):
        return ""
    return clean(row[idx])


def skill_level(value):
    text = clean(value)
    if not text:
        return 0
    digits = "".join(ch for ch in text if ch.isdigit())
    if not digits:
        return 0
    return max(0, min(5, int(digits)))


songs = []
for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=1):
    category = row_value(row, "분류")
    title = row_value(row, "노래 제목", "제목")
    artist = row_value(row, "아티스트", "가수")
    memo = row_value(row, "메모", "비고")
    inst_url = row_value(row, "Inst", "Inst 링크", "인스트", "MR", "MR 링크")
    jeongwa_clip_url = row_value(row, "정와클립", "정와 클립", "클립", "클립 링크")
    skill = skill_level(row_value(row, "숙련도", "숙련"))

    if not (category or title or artist or memo or inst_url or jeongwa_clip_url or skill):
        continue
    songs.append(
        {
            "id": idx,
            "category": category,
            "title": title,
            "artist": artist,
            "instUrl": inst_url,
            "jeongwaClipUrl": jeongwa_clip_url,
            "skillLevel": skill,
            "memo": memo,
        }
    )

payload = "window.JEONGWA_SONGS = "
payload += json.dumps(songs, ensure_ascii=False, indent=2)
payload += ";\n"

OUTPUT.write_text(payload, encoding="utf-8")
print(f"Wrote {len(songs)} songs to {OUTPUT}")
